# app/services/certificate_excel_export.py
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter


class CertificateExcelExporter:

    COLUMNS = [
        ("№", "index"),
        ("Серия", "series"),
        ("Номер", "number"),
        ("Дата создания", "create_date"),
        ("Дата выдачи", "issue_date"),
        ("Причина", "reason"),
        ("Клиент", "client"),
        ("Номинал", "total_amount"),
        ("Остаток", "balance"),
        ("Место выдачи", "issue_place"),
        ("Группа услуг", "servicegroup"),
        ("Создал", "creator"),
        ("Изменил", "editor"),
        ("Примечание", "note"),
    ]

    def export(self, certificates):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Сертификаты"

        # Заголовки
        for column_index, (title, _) in enumerate(
            self.COLUMNS,
            start=1
        ):
            cell = sheet.cell(
                row=1,
                column=column_index,
                value=title
            )
            cell.font = Font(bold=True)
            cell.alignment = Alignment(
                horizontal="center"
            )

        # Данные
        for row_index, certificate in enumerate(
            certificates,
            start=2
        ):
            values = self._get_values(
                certificate,
                row_index - 1
            )

            for column_index, (_, key) in enumerate(
                self.COLUMNS,
                start=1
            ):
                value = values.get(key, "")

                cell = sheet.cell(
                    row=row_index,
                    column=column_index,
                    value=value
                )

                if key in ("total_amount", "balance"):
                    cell.number_format = '# ##0.00'

                elif key in ("create_date", "issue_date"):
                    cell.number_format = 'DD-MM-YYYY'

        # Ширина колонок
        widths = {
            "№": 8,
            "Серия": 15,
            "Номер": 15,
            "Дата создания": 16,
            "Дата выдачи": 16,
            "Причина": 25,
            "Клиент": 30,
            "Номинал": 15,
            "Остаток": 15,
            "Место выдачи": 25,
            "Группа услуг": 25,
            "Создал": 20,
            "Изменил": 20,
            "Примечание": 40,
        }

        for column_index, (title, _) in enumerate(
            self.COLUMNS,
            start=1
        ):
            sheet.column_dimensions[
                get_column_letter(column_index)
            ].width = widths.get(title, 18)

        # Автофильтр Excel
        sheet.auto_filter.ref = sheet.dimensions

        # Закрепляем заголовок
        sheet.freeze_panes = "A2"

        output = BytesIO()
        workbook.save(output)
        output.seek(0)

        return output

    def _get_values(self, certificate, index):
        return {
            "index": index,
            "series": certificate.series,
            "number": (
                certificate.number
                if certificate.number is not None
                else "не присвоен"
            ),
            "create_date": certificate.create_date,
            "issue_date": certificate.issue_date,
            "reason": certificate.reason or "",
            "client": (
                certificate.holder.full_name
                if certificate.holder
                else ""
            ),
            "total_amount": certificate.total_amount,
            "balance": certificate.balance,
            "issue_place": (
                certificate.issuing_place.name
                if certificate.issuing_place
                else "Не указано"
            ),
            "servicegroup": (
                certificate.servicegroup.name
                if certificate.servicegroup
                else "Не указано"
            ),
            "creator": (
                certificate.creator.name
                if certificate.creator
                else ""
            ),
            "editor": (
                certificate.editor.name
                if certificate.editor
                else ""
            ),
            "note": certificate.note or "",
        }