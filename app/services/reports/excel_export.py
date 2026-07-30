# /app/services/reports/excel_export.py

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


class ExcelReportExporter:

    def export(self, report):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = report.title[:31]
        row_index = 1

        # Заголовок отчета
        sheet.cell(
            row=row_index,
            column=1,
            value=report.title
        )
        sheet.cell(
            row=row_index,
            column=1
        ).font = Font(
            bold=True,
            size=14
        )
        row_index += 2

        # Сводка

        for item in report.summary:
            sheet.cell(
                row=row_index,
                column=1,
                value=item.label
            )
            sheet.cell(
                row=row_index,
                column=2,
                value=item.value
            )
            row_index += 1

        row_index += 1

        # Заголовки таблицы
        for column_index, column in enumerate(
            report.columns,
            start=1
        ):
            cell = sheet.cell(
                row=row_index,
                column=column_index,
                value=column.title
            )
            cell.font = Font(
                bold=True
            )
            cell.alignment = Alignment(
                horizontal="center"
            )
        row_index += 1

        # Данные
        for report_row in report.rows:
            self._write_row(
                sheet,
                row_index,
                report.columns,
                report_row
            )
            row_index += 1
            
            # дочерние строки
            for child in report_row.children:
                self._write_row(
                    sheet,
                    row_index,
                    report.columns,
                    child
                )
                row_index += 1

        # ширина колонок
        for column_index, column in enumerate(
            report.columns,
            start=1
        ):
            sheet.column_dimensions[
                get_column_letter(column_index)
            ].width = 18

        output = BytesIO()
        workbook.save(output)
        output.seek(0)

        return output



    def _write_row(
        self,
        sheet,
        row_index,
        columns,
        row
    ):

        if row.row_type == "detail":
            values = {
                "certificate": row.values.get("service"),
                "spent": row.values.get("amount"),
            }

        else:
            values = row.values
            
        for column_index, column in enumerate(columns, start=1):
            value = values.get(column.key, "")

            cell = sheet.cell(
                row=row_index,
                column=column_index,
                value=value
            )

            if column.data_type == "money":
                cell.number_format = '# ##0.00'

            elif column.data_type == "number":
                cell.number_format = '# ##0'

            elif column.data_type == "date":
                cell.number_format = 'DD-MM-YYYY'